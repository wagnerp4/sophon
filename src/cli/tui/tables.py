from __future__ import annotations

import csv
import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
_CELL_REF_RE = re.compile(r"^([A-Z]+)(\d+)$")
_MAX_ROWS = 500
_MAX_COLS = 40
_CELL_CHARS = 64
# TODO: sheet picker for multi-sheet xlsx
# TODO: parquet / ods / xls readers
# TODO: virtual scrolling past _MAX_ROWS


@dataclass
class TableData:
    headers: list[str]
    rows: list[list[str]]
    sheet: str = ""
    truncated: bool = False
    extra_sheets: list[str] = field(default_factory=list)


def load_table(path: Path, text: str | None = None, raw: bytes | None = None) -> TableData:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".tab"}:
        if text is None:
            text = (raw or b"").decode("utf-8-sig", errors="replace")
        delim = "\t" if suffix in {".tsv", ".tab"} else ","
        return load_delimited(text, delim)
    if suffix == ".jsonl":
        if text is None:
            text = (raw or b"").decode("utf-8", errors="replace")
        return load_jsonl(text)
    if suffix == ".xlsx":
        data = raw if raw is not None else path.read_bytes()
        return load_xlsx(data)
    raise ValueError(f"Unsupported table type: {suffix or path.name}")


def load_delimited(text: str, delimiter: str) -> TableData:
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    records = [list(row) for row in reader]
    return _from_records(records)


def load_jsonl(text: str) -> TableData:
    objects: list[dict] = []
    keys: list[str] = []
    seen: set[str] = set()
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        item = json.loads(stripped)
        if not isinstance(item, dict):
            raise ValueError("JSONL preview expects one object per line.")
        objects.append(item)
        for key in item:
            if key not in seen:
                seen.add(key)
                keys.append(key)
    rows = [[_fmt_cell(obj.get(key, "")) for key in keys] for obj in objects]
    return _trim(TableData(headers=keys or ["value"], rows=rows))


def load_xlsx(data: bytes) -> TableData:
    try:
        import openpyxl
    except ImportError:
        return _load_xlsx_zip(data)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    names = list(workbook.sheetnames)
    if not names:
        raise ValueError("Workbook has no sheets.")
    sheet = workbook[names[0]]
    records: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        records.append([_fmt_cell(cell) for cell in row])
        if len(records) > _MAX_ROWS:
            break
    workbook.close()
    table = _from_records(records)
    table.sheet = names[0]
    table.extra_sheets = names[1:]
    return table


def build_xlsx(headers: list[str], rows: list[list[str]], sheet: str = "Sheet1") -> bytes:
    shared = list(headers)
    for row in rows:
        shared.extend(row)
    index = {value: i for i, value in enumerate(shared)}
    sheet_name = _xml_escape(sheet[:31] or "Sheet1")
    header_xml = []
    for col, value in enumerate(headers, start=1):
        header_xml.append(_xlsx_cell(1, col, index[value]))
    row_xml = [f'<row r="1">{"".join(header_xml)}</row>']
    for r_i, row in enumerate(rows, start=2):
        cells = []
        for c_i, value in enumerate(row, start=1):
            cells.append(_xlsx_cell(r_i, c_i, index[value]))
        row_xml.append(f'<row r="{r_i}">{"".join(cells)}</row>')
    sst_items = "".join(f"<si><t>{_xml_escape(item)}</t></si>" for item in shared)
    sheet_body = "".join(row_xml)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            "</Types>",
        )
        zf.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>",
        )
        zf.writestr(
            "xl/_rels/workbook.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
            "</Relationships>",
        )
        zf.writestr(
            "xl/workbook.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<sheets><sheet name="{sheet_name}" sheetId="1" r:id="rId1"/></sheets></workbook>',
        )
        zf.writestr(
            "xl/sharedStrings.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'count="{len(shared)}" uniqueCount="{len(shared)}">{sst_items}</sst>',
        )
        zf.writestr(
            "xl/worksheets/sheet1.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f"<sheetData>{sheet_body}</sheetData></worksheet>",
        )
    return buf.getvalue()


def _from_records(records: list[list[str]]) -> TableData:
    if not records:
        return TableData(headers=["(empty)"], rows=[])
    width = max(len(row) for row in records)
    padded = [row + [""] * (width - len(row)) for row in records]
    headers = [_fmt_cell(value) or f"col_{i + 1}" for i, value in enumerate(padded[0])]
    rows = [[_fmt_cell(value) for value in row] for row in padded[1:]]
    return _trim(TableData(headers=headers, rows=rows))


def _trim(table: TableData) -> TableData:
    headers = table.headers[:_MAX_COLS]
    rows = [row[:_MAX_COLS] for row in table.rows[:_MAX_ROWS]]
    truncated = len(table.headers) > _MAX_COLS or len(table.rows) > _MAX_ROWS
    return TableData(
        headers=headers or ["(empty)"],
        rows=rows,
        sheet=table.sheet,
        truncated=truncated,
        extra_sheets=table.extra_sheets,
    )


def _fmt_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = " ".join(text.splitlines())
    if len(text) > _CELL_CHARS:
        return text[: _CELL_CHARS - 1] + "…"
    return text


def _load_xlsx_zip(data: bytes) -> TableData:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = _xlsx_sheet_names(zf)
        shared = _xlsx_shared_strings(zf)
        xml = zf.read("xl/worksheets/sheet1.xml")
    records = _xlsx_sheet_records(xml, shared)
    table = _from_records(records)
    if names:
        table.sheet = names[0]
        table.extra_sheets = names[1:]
    return table


def _xlsx_sheet_names(zf: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zf.read("xl/workbook.xml"))
    except KeyError:
        return []
    names: list[str] = []
    for sheet in root.findall("m:sheets/m:sheet", _NS):
        name = sheet.attrib.get("name")
        if name:
            names.append(name)
    return names


def _xlsx_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    values: list[str] = []
    for si in root.findall("m:si", _NS):
        texts = [node.text or "" for node in si.iter() if node.tag.endswith("}t") or node.tag == "t"]
        values.append("".join(texts))
    return values


def _xlsx_sheet_records(xml: bytes, shared: list[str]) -> list[list[str]]:
    root = ET.fromstring(xml)
    grid: dict[tuple[int, int], str] = {}
    max_row = -1
    max_col = -1
    for cell in root.findall(".//m:c", _NS):
        ref = cell.attrib.get("r")
        if not ref:
            continue
        parsed = _parse_cell_ref(ref)
        if parsed is None:
            continue
        col, row = parsed
        grid[(row, col)] = _xlsx_cell_value(cell, shared)
        max_row = max(max_row, row)
        max_col = max(max_col, col)
    if max_row < 0:
        return []
    records: list[list[str]] = []
    for row in range(max_row + 1):
        records.append([grid.get((row, col), "") for col in range(max_col + 1)])
    return records


def _xlsx_cell_value(cell: ET.Element, shared: list[str]) -> str:
    kind = cell.attrib.get("t", "n")
    if kind == "s":
        raw = cell.findtext("m:v", default="", namespaces=_NS)
        try:
            return shared[int(raw)]
        except (ValueError, IndexError):
            return raw
    if kind == "inlineStr":
        texts = [node.text or "" for node in cell.iter() if node.tag.endswith("}t") or node.tag == "t"]
        return "".join(texts)
    if kind == "b":
        raw = cell.findtext("m:v", default="", namespaces=_NS)
        return "TRUE" if raw in {"1", "true", "TRUE"} else "FALSE"
    return cell.findtext("m:v", default="", namespaces=_NS)


def _parse_cell_ref(ref: str) -> tuple[int, int] | None:
    match = _CELL_REF_RE.match(ref.upper())
    if match is None:
        return None
    col = 0
    for ch in match.group(1):
        col = col * 26 + (ord(ch) - 64)
    return col - 1, int(match.group(2)) - 1


def _xlsx_cell(row: int, col: int, sst_index: int) -> str:
    return f'<c r="{_col_name(col)}{row}" t="s"><v>{sst_index}</v></c>'


def _col_name(col: int) -> str:
    n = col
    letters = []
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters.append(chr(65 + rem))
    return "".join(reversed(letters)) or "A"


def _xml_escape(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )
